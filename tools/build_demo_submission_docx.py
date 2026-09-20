"""Build the current game's AI Demo submission brief as a Word document."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "01_洞见_AI原生游戏Demo申报说明.docx"
SCREEN_1 = ROOT / "Build" / "Game" / "codex_level1_stone_portal_v1.png"
INK = RGBColor(26, 33, 42)
MUTED = RGBColor(88, 98, 108)
PALE = "F2F5F7"
BORDER = "D9D9D9"


def set_font(style, size, bold=False, color=INK):
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), "Microsoft YaHei")


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def border_cell(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        line = OxmlElement(f"w:{side}")
        line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), "4")
        line.set(qn("w:color"), BORDER)
        borders.append(line)


def pad_cell(cell, top=105, start=125, bottom=105, end=125):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, amount in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(amount))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcpr.append(mar)


def table(doc, headers, rows, widths):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.autofit = False
    for idx, width in enumerate(widths):
        tbl.columns[idx].width = Inches(width)
    for idx, header in enumerate(headers):
        c = tbl.rows[0].cells[idx]
        c.text = header
        shade(c, "E7EDF1")
        c.width = Inches(widths[idx])
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        border_cell(c)
        pad_cell(c)
        for run in c.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)
    for row_idx, values in enumerate(rows):
        cells = tbl.add_row().cells
        for col_idx, value in enumerate(values):
            c = cells[col_idx]
            c.width = Inches(widths[col_idx])
            c.text = value
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx % 2:
                shade(c, PALE)
            border_cell(c)
            pad_cell(c)
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.12
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return tbl


def paragraph(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def bullet(doc, text):
    p = paragraph(doc, text, "List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    return p


def image(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(path), width=Inches(4.8))
    c = doc.add_paragraph(caption)
    c.style = doc.styles["Caption"]
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.68)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.77)
    sec.right_margin = Inches(0.77)

    normal = doc.styles["Normal"]
    set_font(normal, 10.2)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    set_font(doc.styles["Title"], 21, True)
    doc.styles["Title"].paragraph_format.space_after = Pt(13)
    set_font(doc.styles["Heading 1"], 14, True)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(7)
    set_font(doc.styles["Heading 2"], 11, True)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(9)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(5)
    set_font(doc.styles["Caption"], 8.5, False, MUTED)
    doc.styles["Caption"].paragraph_format.space_after = Pt(10)
    set_font(doc.styles["List Bullet"], 10.2)

    doc.add_paragraph("洞见 AI 原生游戏 Demo 申报说明", style="Title")
    p = paragraph(doc, "项目类别  01 AI 原生游戏 Demo    |    版本  Windows 可运行原型    |    日期  2026 年 9 月 17 日")
    p.paragraph_format.space_after = Pt(10)
    for run in p.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED

    paragraph(doc, "《洞见》是一款以重庆历史空间为舞台的 2D 横版剧情冒险游戏。玩家扮演赶赴洪崖洞打卡的现代青年陈默，因嫌传单阿姨烦而抄近路，在古墙附近卷入不同时代的残影。游戏用追逐、工序节奏、护送货物和防空洞引导四种玩法，让玩家亲自经历“走过别人的路”，最后返回当代作出拍照选择。")
    paragraph(doc, "本说明供评审快速核对原型内容、AI 参与证据和核心提交材料。现有版本已导出 Windows 游戏；5 分钟功能演示视频仍需依照文末脚本录制。")
    image(doc, SCREEN_1, "实机画面  第一关古墙石扣触发后出现时空门")

    doc.add_heading("作品定位与运行方式", level=1)
    paragraph(doc, "原型基于 Godot 4.7.2 与 GDScript，采用 640 × 360 像素横版舞台和远中近景分层视差。A / D 移动、Space 跳跃、E 互动；鼠标用于揭开漫画和拖动引导精灵。Windows 试玩包位于 Build/洞见-试玩包-20260917.zip，exe 与 pck 解压到同一目录即可运行。")

    doc.add_page_break()
    doc.add_heading("核心玩法与叙事链", level=1)
    paragraph(doc, "五关以同一角色的视角连续推进；逐格漫画和关间插画交代事件原因，玩家完成当前人物托付的任务后才进入下一段历史。")
    table(doc, ["关卡", "玩家目标与可验证机制"], [
        ("第一关 现代洪崖洞", "传单阿姨追逐，流动人群和障碍阻挡前进；跳跃避让后触发古墙石扣，传送门才出现。"),
        ("第二关 古代窑场", "帮助匠人开壁泄洪；在移动条目标区按键，依次完成烤热、淬水、凿击，共三轮九次，计时内完成。"),
        ("第三关 古镇运货", "替掌柜送染布，沿途对话选择影响货物完整度；即使受损仍须送到渔夫处验货，非完好状态才判失败。"),
        ("第四关 近代防空洞", "带领六名群众穿越洞道；导弹预警后落下，可借挡板防护。轰炸逐轮加密，队伍中有人被击中即失败。"),
        ("第五关 当代归来", "重访洪崖洞并选择是否拍照；拍照路线保存拍摄当刻的本机年月日与时间，进入对应结局画面。"),
    ], [1.50, 5.38])
    paragraph(doc, "贯穿全程的“渝灯”精灵位于右上角，可拖动。它会在进入关卡后自动讲解目标和时代背景，并对点击、拖动做简短回应；正式人物对白出现时让出界面。")

    doc.add_heading("AI 如何进入游戏内容", level=1)
    paragraph(doc, "本项目的 AI 应用主要发生在内容生产与迭代阶段，而非运行时调用大模型。AI 生成的角色动作、互动道具和漫画画格经过透明化、拆帧、接地校正与 Godot 实机验收，成为游戏正式内容。AI 还参与玩法脚本及回归测试。")
    table(doc, ["内容环节", "已落地的实例", "可核查材料"], [
        ("角色与引导", "陈默待机、行走、跳跃及交互帧；第五关传单阿姨；渝灯四帧悬浮形象", "assets/production/hero、npc、companion；asset-manifest.json"),
        ("玩法物件", "古墙、木头、壶、货包状态、导弹与三态挡板等美术与碰撞节点配套", "assets/production/props/gameplay；各关场景与脚本"),
        ("叙事画面", "序章四格和关间漫画交代追逐、凿壁、运货与护送的动机", "assets/production/story；StoryComic 场景资源"),
    ], [1.18, 3.52, 2.18])
    paragraph(doc, "界限说明：当前没有实时 AI 推理、动态关卡或机器学习自适应难度。移动条、轰炸和群众跟随均由脚本控制；若规则要求运行时 AI，还需另行开发。")

    doc.add_page_break()
    doc.add_heading("技术与美术资料", level=1)
    doc.add_heading("制作工具与职责", level=2)
    table(doc, ["工具或系统", "实际用途"], [
        ("DeepSeek Harness 与 Codex", "参与早期场景搭建、后续玩法与剧情迭代、GDScript 调试及自动化回归；均为开发阶段工具。"),
        ("OpenAI image_gen", "生成陈默、渝灯、第五关阿姨及重要互动道具等正式美术；提示词和资产状态写入清单。"),
        ("generate2dsprite 与素材处理脚本", "拆帧、透明化、尺寸与脚底基准统一，并检查空帧和裁切。"),
        ("Godot 4.7.2", "整合场景、动画、视差、对白、音频、物理与 Windows 导出；不是 AI 推理引擎。"),
    ], [2.12, 4.76])
    doc.add_heading("资源来源与使用边界", level=2)
    bullet(doc, "美术：正式 AI 生成资源、提示词记录及验收状态见 assets/production/asset-manifest.json；整体风格规范见 docs/art-direction-brief.md。部分早期素材库被标为“AI 生成非商用”，对外发布前应逐项复核使用范围。")
    bullet(doc, "声音：最新关卡配乐与主要环境音由项目提供者提供并转为 Ogg；来源映射见 assets/audio/user_music/README.md 与 user_ambience/README.md。辅助反馈音含 Kenney CC0 资源。")
    bullet(doc, "视频与字体：开场影片由项目提供者的视频转换并与新主界面配乐叠加，原音轨保留；视频处理记录见 assets/video/SOURCES.md。中文像素字体为 Fusion Pixel Font，许可见 assets/README.md。")

    doc.add_page_break()
    doc.add_heading("核心提交材料与演示计划", level=1)
    table(doc, ["材料", "当前状态", "提交前动作"], [
        ("可运行游戏程序", "已完成 Windows exe 与 pck；另有试玩压缩包", "在目标电脑解压启动，确认视频、音频与第四关六人跟随"),
        ("5 分钟内功能演示视频", "尚未录制为单独提交文件", "按下方脚本录屏，控制总时长不超过 5 分钟"),
        ("核心玩法文档", "本文件已覆盖五关目标、输入与机制", "按主办方模板调整格式和申报信息"),
        ("AI 工具与美术资源说明", "本文件列出工具、内容实例和来源路径", "补齐对外发布所需的素材与音乐授权证明"),
    ], [1.44, 2.26, 3.18])
    doc.add_heading("五分钟视频建议镜头", level=2)
    for text in (
        "00:00 至 00:25 片头短剪、标题和主界面，展示可从游戏程序实际启动。",
        "00:25 至 01:15 第一关追逐、人群阻挡、石扣触发与传送门。",
        "01:15 至 02:05 第二关渝灯讲解、移动条命中和石壁变化。",
        "02:05 至 02:50 第三关货包与关键选择，展示交货后的完整度判定。",
        "02:50 至 03:45 第四关六人跟随、预警、挡板拦截及短暂爆炸照明。",
        "03:45 至 04:35 第五关洪崖洞拍照、实时日期与对应结局。",
        "04:35 至 04:55 用项目内资产清单、帧表和提示词记录收束 AI 内容生产证据。",
    ):
        bullet(doc, text)
    paragraph(doc, "录制时建议保留操作和游戏原声，不以素材剪辑代替实机玩法；视频结尾可用简短字幕标明“AI 参与内容生成与开发，当前无运行时模型推理”。")

    doc.core_properties.title = "洞见 AI 原生游戏 Demo 申报说明"
    doc.core_properties.subject = "游戏内容 核心玩法 AI 工具 美术资源 提交材料"
    doc.core_properties.author = "洞见项目组"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
